import shutil

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==============================
# 設定
# ==============================

USAGE_FILE = "input/利用実績.xlsx"
WAGE_FILE = "input/工賃実績.xlsx"

ERROR_OUTPUT = "output/エラー一覧.xlsx"
MONTHLY_OUTPUT = "output/月別集計.xlsx"

KEY_COLUMNS = ["利用者ID", "対象月"]

OFFICE_FILE = "input/事業所情報.xlsx"
CITY_TEMPLATE = "input/名古屋市_工賃実績報告.xlsx"

CITY_OUTPUT = "output/提出用_工賃実績報告.xlsx"


# ==============================
# 共通処理
# ==============================

def format_target_month(series):
    """対象月を「2026年4月」の形式に変換する"""
    return (
        series.dt.year.astype(str)
        + "年"
        + series.dt.month.astype(str)
        + "月"
    )


def format_excel(path, money_columns=None):
    """Excel内の全シートを担当者が読みやすい形式に整える"""
    wb = load_workbook(path)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    thin = Side(
        style="thin",
        color="D9E1F2"
    )

    for ws in wb.worksheets:

        # ------------------------------
        # 見出し
        # ------------------------------
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # ------------------------------
        # 罫線・配置
        # ------------------------------
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(
                    bottom=thin
                )
                cell.alignment = Alignment(
                    vertical="center"
                )

        # ------------------------------
        # 金額表示
        # ------------------------------
        if money_columns:
            headers = {
                cell.value: cell.column
                for cell in ws[1]
            }

            for column_name in money_columns:
                col = headers.get(column_name)

                if col:
                    for row_number in range(
                        2,
                        ws.max_row + 1
                    ):
                        ws.cell(
                            row_number,
                            col
                        ).number_format = '#,##0"円"'

        # ------------------------------
        # 列幅
        # ------------------------------
        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )

            letter = get_column_letter(
                column_cells[0].column
            )

            ws.column_dimensions[
                letter
            ].width = min(
                max_length + 4,
                50
            )

        # ------------------------------
        # 操作性
        # ------------------------------
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(path)


# ==============================
# 1. データ取込
# ==============================

usage_df = pd.read_excel(
    USAGE_FILE,
    sheet_name="利用実績"
)

wage_df = pd.read_excel(
    WAGE_FILE,
    sheet_name="工賃実績"
)

# 支給日から支給月を算出
wage_df["支給月"] = (
    wage_df["支給日"]
    .dt.to_period("M")
    .astype(str)
)

# 対象月を比較用の年月に変換
wage_df["対象月_比較用"] = (
    wage_df["対象月"]
    .dt.to_period("M")
    .astype(str)
)

# 対象月と支給月が異なるデータを確認対象とする
wage_df["対象月支給月確認"] = (
    wage_df["対象月_比較用"]
    != wage_df["支給月"]
)

office_df = pd.read_excel(
    OFFICE_FILE,
    sheet_name="事業所情報"
)


# ==============================
# 2. 工賃実績の重複チェック
# ==============================

duplicate_mask = wage_df.duplicated(
    subset=KEY_COLUMNS,
    keep=False
)

duplicate_wages = wage_df[
    duplicate_mask
].copy()

duplicate_errors = (
    duplicate_wages[
        ["利用者ID", "氏名", "対象月"]
    ]
    .drop_duplicates(
        subset=KEY_COLUMNS
    )
    .copy()
)

duplicate_errors["エラー内容"] = (
    "同じ利用者・対象月の工賃実績が重複しています"
)

duplicate_errors["対応方法"] = (
    "同一利用者・対象月の工賃実績を確認し、"
    "不要な重複データを修正または削除してください"
)


# ==============================
# 3. 利用実績を月単位に集計
# ==============================

usage_monthly = (
    usage_df
    .groupby(
        [
            "利用者ID",
            "氏名",
            "対象月"
        ],
        as_index=False
    )
    .agg(
        利用日数=("利用日", "count")
    )
)


# ==============================
# 4. 利用実績と工賃実績を突合
# ==============================

check_df = usage_monthly.merge(
    wage_df[
        [
            "利用者ID",
            "対象月",
            "工賃額"
        ]
    ],
    on=KEY_COLUMNS,
    how="left",
    indicator=True
)

missing_wages = check_df[
    check_df["_merge"] == "left_only"
].copy()

missing_errors = missing_wages[
    [
        "利用者ID",
        "氏名",
        "対象月"
    ]
].copy()

missing_errors["エラー内容"] = (
    "利用実績がありますが、工賃実績がありません"
)

missing_errors["対応方法"] = (
    "対象月の工賃支給実績を確認し、"
    "未登録の場合は工賃実績を追加してください"
)


# ==============================
# 5. 対象月・支給月の確認
# ==============================

payment_month_checks = wage_df[
    wage_df["対象月支給月確認"]
].copy()

payment_month_checks = payment_month_checks[
    [
        "利用者ID",
        "氏名",
        "対象月",
        "支給月"
    ]
]

payment_month_checks = payment_month_checks.drop_duplicates(
    subset=[
        "利用者ID",
        "対象月",
        "支給月"
    ]
)

payment_month_checks["エラー内容"] = (
    "対象月と支給月が異なります"
)

payment_month_checks["対応方法"] = (
    "遡及支給等の可能性があるため、"
    "自治体の算定ルールと元データを確認してください"
)


# ==============================
# 6. エラー・要確認一覧を作成
# ==============================

error_df = pd.concat(
    [
        duplicate_errors,
        missing_errors,
        payment_month_checks
    ],
    ignore_index=True
)

error_df = error_df[
    [
        "利用者ID",
        "氏名",
        "対象月",
        "支給月",
        "エラー内容",
        "対応方法"
    ]
]

error_df["対象月"] = format_target_month(
    error_df["対象月"]
)


# ==============================
# 7. 正常な工賃データを抽出
# ==============================

valid_wages = wage_df[
    ~duplicate_mask
].copy()


# ==============================
# 8. 月別集計を作成
# ==============================

monthly_report = (
    valid_wages
    .groupby(
        "対象月",
        as_index=False
    )
    .agg(
        利用者数=("利用者ID", "nunique"),
        工賃支給総額=("工賃額", "sum")
    )
)

# 行政上の正式な「工賃平均額」とは区別した参考値
monthly_report["参考：利用者あたり工賃額"] = (
    monthly_report["工賃支給総額"]
    / monthly_report["利用者数"]
)

monthly_report["対象月"] = format_target_month(
    monthly_report["対象月"]
)


# ==============================
# 9. 集計明細を作成
# ==============================

detail_report = valid_wages[
    [
        "利用者ID",
        "氏名",
        "対象月",
        "支給日",
        "工賃額",
        "事業所ID"
    ]
].copy()

detail_report["対象月"] = format_target_month(
    detail_report["対象月"]
)


# ==============================
# 10. 元データを作成
# ==============================

source_report = wage_df.copy()

source_report["対象月"] = format_target_month(
    source_report["対象月"]
)


# ==============================
# 11. 行政報告用の年間集計
# ==============================

office = office_df.iloc[0]

capacity = int(office["定員"])
opening_days = int(office["年間開所日数"])
opening_months = int(office["年間開所月数"])

# 重複データは行政集計から除外
annual_wage_total = valid_wages["工賃額"].sum()

# サンプルでは1行 = 1利用者 × 1利用日として集計
annual_usage_total = len(usage_df)


# ==============================
# 12. エラー・要確認一覧を出力
# ==============================

error_df.to_excel(
    ERROR_OUTPUT,
    index=False
)


# ==============================
# 13. 月別集計・明細・元データを出力
# ==============================

with pd.ExcelWriter(
    MONTHLY_OUTPUT,
    engine="openpyxl"
) as writer:

    monthly_report.to_excel(
        writer,
        sheet_name="月別集計",
        index=False
    )

    detail_report.to_excel(
        writer,
        sheet_name="集計明細",
        index=False
    )

    source_report.to_excel(
        writer,
        sheet_name="元データ",
        index=False
    )


# ==============================
# 14. 帳票の見た目を整える
# ==============================

format_excel(
    ERROR_OUTPUT
)

format_excel(
    MONTHLY_OUTPUT,
    money_columns=[
        "工賃支給総額",
        "参考：利用者あたり工賃額",
        "工賃額"
    ]
)


# ==============================
# 15. 名古屋市公式Excelへ転記
# ==============================

# 公式テンプレートをコピー
shutil.copyfile(
    CITY_TEMPLATE,
    CITY_OUTPUT
)

city_wb = load_workbook(
    CITY_OUTPUT
)

city_ws = city_wb["(入力してください）B型"]


# ==============================
# 行政指定セルへ値を転記
# ==============================

# B8：定員
city_ws["B8"] = capacity

# C8：年間工賃支払総額
city_ws["C8"] = annual_wage_total

# D8：年間利用者延べ人数
city_ws["D8"] = annual_usage_total

# E8：年間開所日数
city_ws["E8"] = opening_days

# G8：年間開所月数
city_ws["G8"] = opening_months


# ==============================
# 保存
# ==============================

city_wb.save(
    CITY_OUTPUT
)


# ==============================
# 完了
# ==============================

print("\n処理が完了しました")

print("\n=== 集計結果 ===")
print(f"年間工賃支払総額: {annual_wage_total:,.0f}円")
print(f"年間利用者延べ人数: {annual_usage_total:,}人")
print(f"年間開所日数: {opening_days:,}日")
print(f"年間開所月数: {opening_months}か月")

print("\n=== 出力ファイル ===")
print(f"エラー・要確認一覧: {ERROR_OUTPUT}")
print(f"月別集計: {MONTHLY_OUTPUT}")
print(f"行政提出用: {CITY_OUTPUT}")